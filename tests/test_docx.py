"""Tests for the Word generator (REQ-6, REQ-9)."""

import io
import json
from pathlib import Path

from docx import Document
from docx.shared import Mm

from cv_gen.design import load_design
from cv_gen.generators import RenderRequest, get_generator
from cv_gen.generators.docx import docx_text, length
from cv_gen.model import Resume

ROOT = Path(__file__).parent.parent
SAMPLE = Resume.load(ROOT / "examples" / "sample-resume.yaml")


def render(resume: Resume = SAMPLE, request: RenderRequest | None = None):
    """Render a resume to Word and open the result with python-docx."""
    content = get_generator("docx").render(resume, request or RenderRequest())
    return Document(io.BytesIO(content)), content


def paragraphs_in(document, style: str) -> list[str]:
    """Return the text of every paragraph in a style."""
    return [p.text for p in document.paragraphs if p.style.name == style]


def test_structure_is_ats_friendly() -> None:
    """AC 6.2: Title, Heading 1 and Heading 2 styles; no tables; core properties."""
    document, _ = render()
    assert paragraphs_in(document, "Title") == ["Alex Rivera"]
    assert paragraphs_in(document, "Heading 1") == [
        "Summary",
        "Experience",
        "Projects",
        "Skills",
        "Education",
        "Languages",
        "Talks",
    ]
    headings = paragraphs_in(document, "Heading 2")
    assert headings[:3] == ["Northwind Freight", "Ledgerline", "Freelance"]
    assert document.tables == []
    assert document.core_properties.author == "Alex Rivera"
    assert document.core_properties.title == "Alex Rivera – CV"
    assert "Kafka" in document.core_properties.keywords
    bullets = paragraphs_in(document, "List Bullet")
    assert "Cut p99 latency of the tracking API from 1.2 s to 140 ms." in bullets


def test_design_fonts_and_page(tmp_path: Path) -> None:
    """AC 9.6: the design's body font sets the Normal style; page size and accent apply."""
    tokens = tmp_path / "t.json"
    tokens.write_text(
        json.dumps(
            {
                "font": {"$type": "fontFamily", "body": {"$value": ["Georgia", "serif"]}},
                "color": {"accent": {"$type": "color", "$value": "#ff2bd6"}},
                "page": {
                    "size": {"$value": "A4"},
                    "margin": {"$type": "dimension", "$value": "18mm"},
                },
            }
        ),
        encoding="utf-8",
    )
    document, _ = render(request=RenderRequest(theme=load_design(str(tokens))))
    assert document.styles["Normal"].font.name == "Georgia"
    assert str(document.styles["Heading 1"].font.color.rgb) == "FF2BD6"
    section = document.sections[0]
    assert abs(section.page_width - Mm(210)) < Mm(1)
    assert abs(section.left_margin - Mm(18)) < Mm(1)


def test_links_and_text_extraction() -> None:
    """Contact links are real hyperlinks, and docx_text sees their text."""
    document, content = render()
    rels = [r.target_ref for r in document.part.rels.values() if r.reltype.endswith("/hyperlink")]
    assert "https://github.com/alexrivera-example" in rels
    assert "mailto:alex.rivera@example.com" in rels
    text = docx_text(content)
    assert "alexrivera-example" in text and "Northwind Freight" in text


def test_empty_sections_and_every_layout() -> None:
    """No Education heading without education; every built-in section renders."""
    document, _ = render(SAMPLE.model_copy(update={"education": []}))
    assert "Education" not in paragraphs_in(document, "Heading 1")
    sample = Resume.load(ROOT / "tests" / "fixtures" / "jsonresume-sample.json")
    document, content = render(sample)
    assert len(paragraphs_in(document, "Heading 1")) == 12
    assert "Erlich Bachman" in docx_text(content)


def test_length_conversion() -> None:
    """CSS lengths convert to Word lengths; relative units use the fallback."""
    from docx.shared import Pt

    assert length("12pt", Pt(1)) == Pt(12)
    assert length("16px", Pt(1)) == Pt(12)
    assert length("1pc", Pt(1)) == Pt(12)
    assert length("1.5rem", Pt(7)) == Pt(7)
    assert length("nonsense", Pt(7)) == Pt(7)


def test_core_properties_fit_word_limits() -> None:
    """Long skill lists and names are trimmed to Word's 255-character property limit."""
    data = SAMPLE.to_data()
    data["basics"]["name"] = "N" * 300
    data["skills"] = [
        {"name": f"Group {i}", "keywords": [f"skill-{i}-{j}" for j in range(8)]} for i in range(20)
    ]
    document, _ = render(Resume.from_data(data))
    props = document.core_properties
    words = {w for group in data["skills"] for w in [group["name"], *group["keywords"]]}
    assert 200 < len(props.keywords) <= 255
    assert set(props.keywords.split(", ")) <= words, "no keyword is cut in half"
    assert len(props.title) == 255 and props.title.endswith("…")
