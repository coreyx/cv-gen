"""Word output, built for people and for applicant-tracking systems alike.

ATS parsers read Word structure, not appearance, so the document uses Word's own
``Title``, ``Heading 1`` and ``Heading 2`` styles for the name, sections and entries,
real bulleted lists, no tables or text boxes, and core properties naming the author.
The design system supplies the fonts, sizes, margins and accent color.
"""

from __future__ import annotations

import io
import re

from docx import Document
from docx.document import Document as DocumentType
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Mm, Pt, RGBColor
from docx.text.paragraph import Paragraph

from cv_gen.design.theme import Color, Theme
from cv_gen.generators import view
from cv_gen.generators.base import Generator, RenderRequest
from cv_gen.generators.view import Entry, ViewSection
from cv_gen.model import Resume

PAGE_SIZES = {
    "Letter": (Inches(8.5), Inches(11)),
    "Legal": (Inches(8.5), Inches(14)),
    "A4": (Mm(210), Mm(297)),
    "A5": (Mm(148), Mm(210)),
}


def length(css: str, fallback: Pt) -> Pt | Inches | Mm | Cm:
    """Convert a CSS dimension from the theme to a Word length; relative units fall back."""
    match = re.fullmatch(r"(-?[\d.]+)([a-z%]+)", css)
    if not match:
        return fallback
    number, unit = float(match.group(1)), match.group(2)
    converters = {
        "pt": Pt,
        "px": lambda n: Pt(n * 0.75),
        "in": Inches,
        "mm": Mm,
        "cm": Cm,
        "pc": lambda n: Pt(n * 12),
    }
    return converters[unit](number) if unit in converters else fallback


# Word's core properties (title, subject, keywords) hold at most 255 characters.
CORE_PROPERTY_LIMIT = 255


def _fit(text: str) -> str:
    """Shorten text to fit a core property, ending with an ellipsis when cut."""
    if len(text) <= CORE_PROPERTY_LIMIT:
        return text
    return text[: CORE_PROPERTY_LIMIT - 1].rstrip() + "…"


def _fit_list(items: list[str]) -> str:
    """Join as many whole items as fit in a core property, so no keyword is cut in half."""
    joined = ""
    for item in items:
        candidate = f"{joined}, {item}" if joined else item
        if len(candidate) > CORE_PROPERTY_LIMIT:
            break
        joined = candidate
    return joined


def rgb(color: Color) -> RGBColor:
    """Convert a theme color to Word's RGB (Word has no alpha)."""
    return RGBColor(color.red, color.green, color.blue)


def _set_font(
    style_or_run, family: str, size: Pt | None = None, color: Color | None = None
) -> None:
    """Set a style's or run's font family for every script, plus size and color."""
    font = style_or_run.font
    font.name = family
    # python-docx sets only the Latin font slots; East Asian text needs its own slot or
    # Word substitutes a different font for it.
    rpr = style_or_run.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for slot in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(slot), family)
    if size is not None:
        font.size = size
    if color is not None:
        font.color.rgb = rgb(color)


def add_hyperlink(paragraph: Paragraph, text: str, url: str, color: Color) -> None:
    """Append a clickable link; python-docx has no API for one, so build the XML."""
    part = paragraph.part
    rel_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color_el = OxmlElement("w:color")
    color_el.set(qn("w:val"), color.hex[1:7])
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(color_el)
    props.append(underline)
    run.append(props)
    text_el = OxmlElement("w:t")
    text_el.text = text
    text_el.set(qn("xml:space"), "preserve")
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


class _Builder:
    """Writes one CV into a fresh Word document."""

    def __init__(self, resume: Resume, request: RenderRequest) -> None:
        """Prepare a document styled from the request's theme."""
        self.resume = resume
        self.request = request
        self.theme: Theme = request.resolved_theme()
        self.doc: DocumentType = Document()
        self._style_document()

    def _style_document(self) -> None:
        """Apply the theme's page setup, fonts, sizes and colors to Word's styles."""
        theme, doc = self.theme, self.doc
        section = doc.sections[0]
        size_name = theme.value("page.size").split()[0]
        if size_name in PAGE_SIZES:
            section.page_width, section.page_height = PAGE_SIZES[size_name]
        margin = length(theme.value("page.margin"), Inches(0.75))
        section.left_margin = section.right_margin = margin
        section.top_margin = section.bottom_margin = margin

        body_font = theme.fonts["font.body"][0]
        heading_font = theme.fonts["font.heading"][0]
        base = length(theme.value("font-size.base"), Pt(10.5))
        normal = doc.styles["Normal"]
        _set_font(normal, body_font, base, theme.colors["color.text"])
        normal.paragraph_format.space_after = Pt(3)

        title = doc.styles["Title"]
        _set_font(
            title,
            heading_font,
            length(theme.value("font-size.name"), Pt(22)),
            theme.colors["color.heading"],
        )
        title.font.bold = True
        title.paragraph_format.space_after = Pt(2)
        # Word's default Title has a bottom border; the accent rule belongs to sections.
        ppr = title.element.get_or_add_pPr()
        for border in ppr.findall(qn("w:pBdr")):
            ppr.remove(border)

        h1 = doc.styles["Heading 1"]
        _set_font(
            h1,
            heading_font,
            length(theme.value("font-size.section"), Pt(12)),
            theme.colors["color.accent"],
        )
        h1.font.bold = True
        h1.paragraph_format.space_before = length(theme.value("space.section"), Pt(14))
        h1.paragraph_format.space_after = Pt(4)

        h2 = doc.styles["Heading 2"]
        _set_font(
            h2,
            heading_font,
            length(theme.value("font-size.entry"), Pt(10.5)),
            theme.colors["color.heading"],
        )
        h2.font.bold = True
        h2.paragraph_format.space_before = length(theme.value("space.entry"), Pt(8))
        h2.paragraph_format.space_after = Pt(1)

        for name in ("List Bullet", "Subtitle"):
            _set_font(doc.styles[name], body_font, None, theme.colors["color.text"])

        props = doc.core_properties
        name = self.resume.basics.name or "Curriculum Vitae"
        props.title = _fit(f"{name} – CV")
        props.author = _fit(name)
        props.subject = _fit(self.resume.basics.label or "")
        skills = [kw for skill in self.resume.skills for kw in [skill.name, *skill.keywords] if kw]
        props.keywords = _fit_list(skills)
        props.comments = "Generated by cv-gen"

    def _muted(self, paragraph: Paragraph, text: str, bold_prefix: str | None = None) -> None:
        """Append a meta line: an optional bold lead, then muted text."""
        if bold_prefix:
            paragraph.add_run(bold_prefix).bold = True
            if text:
                paragraph.add_run(" · ")
        if text:
            run = paragraph.add_run(text)
            run.font.color.rgb = rgb(self.theme.colors["color.muted"])

    def header(self) -> None:
        """Write the name, label, headline and contact line."""
        head = view.header(self.resume)
        self.doc.add_paragraph(head.name, style="Title")
        if head.label:
            label = self.doc.add_paragraph()
            run = label.add_run(head.label)
            run.bold = True
            run.font.color.rgb = rgb(self.theme.colors["color.accent"])
        if head.headline:
            self.doc.add_paragraph().add_run(head.headline).italic = True
        if head.contacts:
            line = self.doc.add_paragraph()
            for index, contact in enumerate(head.contacts):
                if index:
                    line.add_run("  ·  ")
                if contact.kind not in ("location", "email", "phone", "url"):
                    line.add_run(f"{contact.kind}: ")
                if contact.url and contact.kind not in ("phone",):
                    add_hyperlink(line, contact.text, contact.url, self.theme.colors["color.link"])
                else:
                    line.add_run(contact.text)

    def entry(self, entry: Entry) -> None:
        """Write one ``entries`` block: heading, meta line, summary, bullets, keywords."""
        heading = self.doc.add_paragraph(entry.title or "", style="Heading 2")
        heading.paragraph_format.keep_with_next = True
        meta = [part for part in (entry.dates, entry.location) if part]
        if entry.subtitle or meta:
            line = self.doc.add_paragraph()
            line.paragraph_format.keep_with_next = True
            self._muted(line, " · ".join(meta), entry.subtitle)
        if entry.url:
            add_hyperlink(
                self.doc.add_paragraph(), entry.url, entry.url, self.theme.colors["color.link"]
            )
        if entry.summary:
            self.doc.add_paragraph(entry.summary)
        for bullet in entry.bullets:
            self.doc.add_paragraph(bullet, style="List Bullet")
        if entry.keywords:
            self.doc.add_paragraph().add_run(", ".join(entry.keywords)).italic = True

    def section(self, section: ViewSection) -> None:
        """Write one section under a ``Heading 1``."""
        self.doc.add_paragraph(
            section.title, style="Heading 1"
        ).paragraph_format.keep_with_next = True
        if section.updated:
            self._muted(self.doc.add_paragraph(), f"Updated {section.updated}")
        for paragraph in section.paragraphs:
            self.doc.add_paragraph(paragraph)
        if section.layout == "entries":
            for entry in section.entries:
                self.entry(entry)
        elif section.layout == "keywords":
            for entry in section.entries:
                line = self.doc.add_paragraph()
                label = entry.title or ""
                if entry.subtitle:
                    label += f" ({entry.subtitle})"
                line.add_run(f"{label}: ").bold = True
                line.add_run(", ".join(entry.keywords))
        else:
            for entry in section.entries:
                line = self.doc.add_paragraph(style="List Bullet")
                if entry.title:
                    line.add_run(entry.title).bold = True
                extra = ", ".join(part for part in (entry.subtitle,) if part)
                if extra:
                    line.add_run(f", {extra}")
                if entry.dates:
                    line.add_run(f" ({entry.dates})")
                if entry.summary:
                    line.add_run(f" — {entry.summary}" if entry.title else entry.summary)
                if entry.url:
                    line.add_run(" ")
                    add_hyperlink(line, entry.url, entry.url, self.theme.colors["color.link"])

    def build(self) -> bytes:
        """Write the whole CV and return the .docx bytes."""
        self.header()
        for section in view.sections(self.resume, self.request.options):
            self.section(section)
        buffer = io.BytesIO()
        self.doc.save(buffer)
        return buffer.getvalue()


def docx_text(content: bytes) -> str:
    """Return the visible text of a .docx (paragraphs and hyperlinks), for the guard."""
    document = Document(io.BytesIO(content))
    lines = []
    for paragraph in document.paragraphs:
        texts = [node.text or "" for node in paragraph._p.iter(qn("w:t"))]
        lines.append("".join(texts))
    props = document.core_properties
    lines += [props.title or "", props.subject or "", props.keywords or ""]
    return "\n".join(lines)


class DocxGenerator(Generator):
    """Microsoft Word (``.docx``)."""

    name = "docx"
    extension = "docx"
    binary = True
    description = "Word, with Title/Heading styles, real lists and no tables (ATS-friendly)."

    def render(self, resume: Resume, request: RenderRequest) -> bytes:
        """Render the CV as a Word document."""
        return _Builder(resume, request).build()
